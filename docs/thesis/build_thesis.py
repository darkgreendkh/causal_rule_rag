"""Assemble the draft and render with the installed Microsoft Word on Windows.

Usage: D:/anaconda/python.exe docs/thesis/build_thesis.py [--render]
Requires python-docx, Pillow, pywin32 and pypdfium2; Pandoc lives on D:.
All generated QA pages remain under .runtime/thesis-qa (never committed).
"""

import argparse
import io
import json
import re
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parents[1]
QA = PROJECT / '.runtime' / 'thesis-qa'
TITLE = '基于规则约束与因果增强的智能审批关键技术研究'
STEM = TITLE + '（初稿）'
FIGURES = {
    '1-1': ('technical-route', '研究技术路线'),
    '3-1': ('dual-graph', '规则约束的双层知识图谱'),
    '3-2': ('rpc', 'RPC 结构一致性计算'),
    '3-3': ('scs', 'SCS 跨层支持与冲突拒绝'),
    '4-1': ('community', '强约束收缩与层次社区组织'),
    '4-2': ('directed-search', '社区筛选与有向多跳检索'),
    '4-3': ('evidence', '完整条件核查与证据组织'),
}


def image_md(number, suffix, caption):
    return f'![图{number} {caption}](figures/fig{number}-{suffix}.png){{width=15cm}}'


def result_figure(number, caption):
    """Use the generated chart when the experiment has run, else the pending placeholder."""
    caption = caption.strip().rstrip('。')
    drawn = sorted((ROOT / 'figures').glob(f'fig6-{number}-*.png'))
    if drawn:
        return f'![图6-{number} {caption}](figures/{drawn[0].name}){{width=14cm}}'
    return (f'![图6-{number} {caption}（待标注）]'
            '(figures/experiment-results-placeholder.png){width=12cm}')


def assemble():
    refs = json.loads((ROOT / 'references.json').read_text('utf-8'))
    local = ROOT / 'references-local.json'
    if local.exists():
        values = json.loads(local.read_text('utf-8'))
        refs.update({v['key']: v for v in values} if isinstance(values, list) else values)
    parts = []
    counts = {}
    for path in sorted((ROOT / 'chapters').glob('*.md')):
        if 'writing-note' in path.name:
            continue
        text = path.read_text('utf-8').split('## 本章引用')[0].strip()
        text = re.sub(r'<!-- REFERENCE .*?-->', '', text, flags=re.S)
        text = text.replace('../figures/', 'figures/')
        text = text.replace('fig5-2-repair.png', 'fig5-2-repair-search.png')
        text = re.sub(r'((?:\|[^\n]*\n)+)\n(表\d+-\d+[^\n]*)',
                      lambda m: m[2] + '\n\n' + m[1], text)
        for number, (suffix, caption) in FIGURES.items():
            replacement = image_md(number, suffix, caption)
            text = re.sub(r'【图' + number + r'占位：[\s\S]*?】', lambda _: replacement, text)
            if number == '1-1':
                text = re.sub(r'^图1-1 研究技术路线图[^\n]*', lambda _: replacement, text, flags=re.M)
        text = re.sub(r'<!-- SCREENSHOT:(\w+) (.*?) -->',
                      lambda m: f'![{m[2]}](screenshots/{m[1]}.png){{width=15cm}}', text)
        text = re.sub(r'^图6-(\d+) ([^\n；]*)[^\n]*', lambda m: result_figure(m[1], m[2]),
                      text, flags=re.M)
        if path.name.startswith('05-'):
            architecture = '![图5-1 系统实现架构](figures/fig5-1-architecture.png)'
            text = text.replace(architecture, '')
            position = text.index('\n## 5.2')
            text = text[:position] + '\n\n' + architecture + '{width=15cm}\n' + text[position:]
        if path.name.startswith('07-'):
            text = text.split('\n# 致谢')[0]
        if not path.name.startswith('00-'):
            body = re.sub(r'```.*?```', '', text, flags=re.S)
            body = re.sub(r'^\s*(?:#|\||!\[|图\d|表\d).*$', '', body, flags=re.M)
            counts[path.name] = len(re.findall(r'[\u3400-\u4dbf\u4e00-\u9fff]', body))
        parts.append(text)
    content = '\n\n'.join(parts)
    keys = list(dict.fromkeys(re.findall(r'\[([A-Za-z][A-Za-z0-9]+)\]', content)))
    missing = set(keys) - set(refs)
    if missing:
        raise ValueError(f'Missing references: {missing}')
    for i, key in enumerate(keys, 1):
        content = content.replace(f'[{key}]', f'[{i}]')
    content += '\n\n# 参考文献\n\n'
    link_labels = {'Chen2025': '原刊元数据', 'Zhou2024': '作者学位辅助信息'}
    content += '\n\n'.join(f'[{i}] {refs[key]["text"]} [{link_labels.get(key, "来源")}]({refs[key]["url"]})。'
                           if refs[key].get('url') else f'[{i}] {refs[key]["text"]}'
                           for i, key in enumerate(keys, 1))
    content += '\n\n# 致谢\n\n（个人信息与致谢内容待作者填写。）\n'
    (ROOT / (STEM + '.md')).write_text(content, 'utf-8')
    audit = {'main_text_hanzi': sum(counts.values()), 'chapters': counts,
             'references': [{'number': i, 'key': k, **refs[k]} for i, k in enumerate(keys, 1)],
             'research_results': 'automatic_layer_complete_annotation_layer_pending',
             'experimental_figures': {
                 'generated': sorted(p.name for p in (ROOT / 'figures').glob('fig6-*.png')),
                 'pending_annotation': ['6-1', '6-2', '6-3', '6-8']},
             'counting': 'Unicode Han characters; excludes headings, tables, figure captions and fenced algorithms'}
    (ROOT / 'document-audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), 'utf-8')
    print(json.dumps({'main_text_hanzi': audit['main_text_hanzi'], 'references': len(keys)}))


def set_font(style, size, chinese='宋体', bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), chinese)


def reference_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.7)
    section.left_margin, section.right_margin = Cm(3), Cm(2.7)
    section.header_distance = section.footer_distance = Cm(1.4)
    normal = doc.styles['Normal']
    set_font(normal, 12)
    pf = normal.paragraph_format
    pf.first_line_indent = Pt(24)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(20)
    pf.space_before = pf.space_after = Pt(0)
    pf.widow_control = True
    for name in ['Body Text', 'First Paragraph']:
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        doc.styles[name].base_style = normal
        doc.styles[name].paragraph_format.first_line_indent = Pt(24)
        doc.styles[name].paragraph_format.space_after = Pt(0)
    for level, size in [(1, 16), (2, 14), (3, 12)]:
        style = doc.styles[f'Heading {level}']
        set_font(style, size, '黑体', True)
        p = style.paragraph_format
        p.first_line_indent = Pt(0)
        p.space_before, p.space_after = Pt(14 if level == 1 else 10), Pt(10 if level == 1 else 6)
        p.line_spacing = 1.3
        p.keep_with_next = True
        p.page_break_before = level == 1
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    for name in ['Caption', 'Image Caption', 'Table Caption']:
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        set_font(style, 10.5)
        p = style.paragraph_format
        p.first_line_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.line_spacing = 1.2
        p.space_before, p.space_after = Pt(5), Pt(8)
    for name in ['Source Code', 'Verbatim Char']:
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH if name == 'Source Code' else WD_STYLE_TYPE.CHARACTER)
        set_font(doc.styles[name], 9, '等线')
        doc.styles[name].font.name = 'Consolas'
        if name == 'Source Code':
            p = doc.styles[name].paragraph_format
            p.first_line_indent = Pt(0)
            p.line_spacing = 1.1
    doc.save(QA / 'reference.docx')


def field(paragraph, instruction):
    run = paragraph.add_run()
    start = OxmlElement('w:fldChar')
    start.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText')
    text.set(qn('xml:space'), 'preserve')
    text.text = instruction
    separator = OxmlElement('w:fldChar')
    separator.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for element in [start, text, separator, end]:
        run._r.append(element)


def postprocess(path):
    doc = Document(path)
    doc.settings.odd_and_even_pages_header_footer = False
    # Pandoc keeps the template accent colour on the linked heading character styles.
    for style in doc.styles.element.findall(qn('w:style')):
        name = style.find(qn('w:name'))
        label = name.get(qn('w:val')) if name is not None else ''
        if not re.match(r'(heading|Heading) ?[1-6]|标题 ?[1-6]', label or ''):
            continue
        rpr = style.find(qn('w:rPr'))
        if rpr is None:
            rpr = OxmlElement('w:rPr')
            style.append(rpr)
        for tag in rpr.findall(qn('w:color')):
            rpr.remove(tag)
        black = OxmlElement('w:color')
        black.set(qn('w:val'), '000000')
        rpr.insert(0, black)
    first = doc.paragraphs[0]
    for text, size, before, after in [
        ('硕士学位论文', 26, 60, 65),
        ('基于规则约束与因果增强的\n智能审批关键技术研究', 22, 0, 24),
        ('初稿', 14, 0, 65),
        ('作者姓名：________________', 12, 0, 18),
        ('学科专业：________________', 12, 0, 18),
        ('指导教师：________________', 12, 0, 18),
        ('培养单位：________________', 12, 0, 18),
        ('完成日期：________________', 12, 0, 0),
    ]:
        p = first.insert_paragraph_before(text)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        p.paragraph_format.line_spacing = 1.5
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(size)
            r.font.bold = size >= 22
            r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '黑体' if size >= 22 else '宋体')
    chapter = next(p for p in doc.paragraphs if re.match(r'第\s*1\s*章', p.text))
    toc_heading = chapter.insert_paragraph_before('目录', 'TOC Heading')
    toc_heading.paragraph_format.page_break_before = True
    toc_heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_heading.paragraph_format.first_line_indent = Pt(0)
    toc = chapter.insert_paragraph_before()
    toc.paragraph_format.first_line_indent = Pt(0)
    field(toc, ' TOC \\o "1-3" \\h \\z \\u ')
    # The cover is unnumbered; Roman front matter and Arabic chapter pages follow.
    for section in doc.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.text = TITLE
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header.paragraph_format.first_line_indent = Pt(0)
        for r in header.runs:
            r.font.size = Pt(9)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.first_line_indent = Pt(0)
        field(footer, ' PAGE ')
    for p in doc.paragraphs:
        if p._p.xpath('.//m:oMathPara'):
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(6)
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if re.match(r'^表\d+-\d+', p.text):
            p.style = doc.styles['Table Caption']
            p.paragraph_format.keep_with_next = True
        if re.match(r'^\[\d+\]', p.text):
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.line_spacing = Pt(17)
            p.paragraph_format.space_after = Pt(5)
            for r in p.runs:
                r.font.size = Pt(10.5)
    for shape in doc.inline_shapes:
        scale = min(1, Cm(15.3) / shape.width, Cm(13.8) / shape.height)
        shape.width, shape.height = int(shape.width * scale), int(shape.height * scale)
    for table in doc.tables:
        if table._tbl.xpath('.//w:drawing'):
            continue
        table.autofit = False
        n = len(table.columns)
        widths = [15.3 / n] * n
        if n >= 4:
            widths = [max(2.5, 15.3 / n * 1.3)] + [0] * (n - 1)
            widths[1:] = [(15.3 - widths[0]) / (n - 1)] * (n - 1)
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        properties = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for side in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
            border = OxmlElement('w:' + side)
            border.set(qn('w:val'), 'single' if side in ['top', 'bottom'] else 'nil')
            border.set(qn('w:sz'), '10')
            border.set(qn('w:color'), '000000')
            borders.append(border)
        properties.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for side, value in [('top', '70'), ('bottom', '70'), ('left', '90'), ('right', '90')]:
            item = OxmlElement('w:' + side)
            item.set(qn('w:w'), value)
            item.set(qn('w:type'), 'dxa')
            margins.append(item)
        properties.append(margins)
        for index, row in enumerate(table.rows):
            if index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = 1
                if index == 0:
                    cell_borders = OxmlElement('w:tcBorders')
                    bottom = OxmlElement('w:bottom')
                    bottom.set(qn('w:val'), 'single')
                    bottom.set(qn('w:sz'), '6')
                    bottom.set(qn('w:color'), '000000')
                    cell_borders.append(bottom)
                    cell._tc.get_or_add_tcPr().append(cell_borders)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                    # Short tables stay together, with their caption above them.
                    p.paragraph_format.keep_with_next = len(table.rows) <= 10 and index < len(table.rows) - 1
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(p.text) > 15 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(10)
                        r.font.bold = index == 0
    doc.core_properties.title = TITLE
    doc.core_properties.subject = '硕士学位论文初稿；科研实验结果待填'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.save(path)


def render(path):
    import win32com.client
    import pypdfium2 as pdfium
    word = win32com.client.DispatchEx('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    pdf = QA / 'thesis.pdf'
    try:
        document = word.Documents.Open(str(path.resolve()))
        # Separate cover, front matter and body; each has its own page numbering.
        for heading in ['第1章', '摘要']:
            finder = document.Content.Duplicate
            finder.Find.Text = heading
            if not finder.Find.Execute():
                raise ValueError(f'Heading not found in Word: {heading}')
            finder.Collapse(1)
            finder.InsertBreak(2)
        for section in document.Sections:
            if section.Index > 1:
                section.PageSetup.DifferentFirstPageHeaderFooter = False
                section.Footers(1).PageNumbers.RestartNumberingAtSection = True
                section.Footers(1).PageNumbers.StartingNumber = 1
                section.Footers(1).PageNumbers.NumberStyle = 2 if section.Index == 2 else 0
        document.Fields.Update()
        for toc in document.TablesOfContents:
            toc.Update()
        document.Repaginate()
        for toc in document.TablesOfContents:
            toc.UpdatePageNumbers()
        document.Save()
        document.ExportAsFixedFormat(str(pdf.resolve()), 17)
        pages = document.ComputeStatistics(2)
        document.Close(False)
    finally:
        word.Quit()
    # Word may add its local user name on save; preserve every layout part byte-for-byte.
    content = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(content, 'w') as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == 'docProps/core.xml':
                data = re.sub(rb'(<cp:lastModifiedBy[^>]*>).*?(</cp:lastModifiedBy>)', rb'\1\2', data)
            target.writestr(entry, data)
    path.write_bytes(content.getvalue())
    reader = pdfium.PdfDocument(str(pdf))
    for i in range(len(reader)):
        page = reader[i]
        bitmap = page.render(scale=1.4)
        bitmap.to_pil().save(QA / f'page-{i + 1:03}.png')
        bitmap.close()
        page.close()
    reader.close()
    print(json.dumps({'word_pages': pages, 'rendered_pdf': str(pdf)}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--render', action='store_true')
    args = parser.parse_args()
    QA.mkdir(parents=True, exist_ok=True)
    assemble()
    reference_docx()
    output = ROOT / (STEM + '.docx')
    pandoc = PROJECT / '.runtime/tools/pandoc-3.11/pandoc.exe'
    subprocess.run([str(pandoc), str(ROOT / (STEM + '.md')), '-o', str(output),
                    '--from=markdown+tex_math_single_backslash+tex_math_dollars+implicit_figures',
                    '--reference-doc=' + str(QA / 'reference.docx'),
                    '--resource-path=' + str(ROOT)], check=True)
    postprocess(output)
    if args.render:
        render(output)
